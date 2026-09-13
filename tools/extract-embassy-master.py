import json
import re
import sys
import urllib.parse
from collections import Counter, defaultdict
from pathlib import Path

import openpyxl


EXPECTED_COUNTS = {
    1: 13,
    2: 14,
    3: 19,
    4: 18,
    5: 16,
    6: 16,
    7: 22,
    8: 21,
    9: 10,
    10: 8,
}

DAY1_IDS = {
    "ペルー": "peru",
    "チェコ": "czech",
    "クロアチア": "croatia",
    "オマーン": "oman",
    "スイス": "switzerland",
    "ボスニア・ヘルツェゴビナ": "bosnia-herzegovina",
    "フランス": "france",
    "パキスタン": "pakistan",
    "キプロス": "cyprus",
    "ドイツ": "germany",
    "フィンランド": "finland",
    "イラン": "iran",
    "韓国": "korea",
}

SOURCE_PATH = Path(
    r"C:\Users\k_ota\OneDrive\Documents\家族共有ファイル\旅行\embassy_stamp_rally_157攻略表_住所マスター確定版_20260906.xlsx"
)
ROOT = Path(__file__).resolve().parents[1]


def normalize_id_text(value):
    text = str(value).strip()
    text = text.replace("・", "-")
    text = text.replace("／", "-")
    text = text.replace("/", "-")
    text = re.sub(r"\s+", "-", text)
    text = re.sub(r"[（）()、。，,.]", "", text)
    text = re.sub(r"-+", "-", text)
    return text.strip("-")


def map_query(value):
    text = str(value or "").strip()
    if not text.startswith("https://www.google.com/maps/search/"):
        return "", text
    parsed = urllib.parse.urlparse(text)
    query = urllib.parse.parse_qs(parsed.query).get("query", [""])[0]
    return query, text


def extract_rows():
    workbook = openpyxl.load_workbook(SOURCE_PATH, data_only=True, read_only=True)
    rows = []
    issues = []

    for day in range(1, 11):
        sheet_name = f"Day{day:02d} 実行表"
        if sheet_name not in workbook.sheetnames:
            issues.append({"type": "missing_sheet", "day": day, "sheet": sheet_name})
            continue
        sheet = workbook[sheet_name]
        for values in sheet.iter_rows(min_row=7, max_row=80, values_only=True):
            if not values[0]:
                continue
            order, country, embassy_name, address, map_value = values[:5]
            query, source_map = map_query(map_value)
            rows.append(
                {
                    "country": str(country).strip(),
                    "embassyName": str(embassy_name).strip(),
                    "day": day,
                    "order": int(order),
                    "address": str(address).strip(),
                    "googleMapsQuery": query,
                    "sourceMap": source_map,
                }
            )
    return rows, issues


def assign_ids(rows):
    used = set()
    for row in rows:
        if row["day"] == 1 and row["country"] in DAY1_IDS:
            row["id"] = DAY1_IDS[row["country"]]
        else:
            row["id"] = f"embassy-{normalize_id_text(row['country'])}"

        if row["id"] in used:
            row["id"] = f"{row['id']}-{normalize_id_text(row['embassyName'])}"
        used.add(row["id"])


def js_string(value):
    return json.dumps(value, ensure_ascii=False)


def write_data_js(rows):
    lines = ["const EMBASSY_MASTER = ["]
    for index, row in enumerate(rows):
        comma = "," if index < len(rows) - 1 else ""
        lines.extend(
            [
                "  {",
                f"    id: {js_string(row['id'])},",
                f"    country: {js_string(row['country'])},",
                f"    embassyName: {js_string(row['embassyName'])},",
                f"    day: {row['day']},",
                f"    order: {row['order']},",
                f"    address: {js_string(row['address'])},",
                f"    googleMapsQuery: {js_string(row['googleMapsQuery'])},",
                f"    sourceMap: {js_string(row['sourceMap'])}",
                f"  }}{comma}",
            ]
        )
    lines.append("];")
    (ROOT / "js" / "data.js").write_text("\n".join(lines) + "\n", encoding="utf-8")


def validate(rows, extraction_issues):
    issues = list(extraction_issues)
    counts = Counter(row["day"] for row in rows)
    ids = [row["id"] for row in rows]
    id_counts = Counter(ids)

    for day, expected in EXPECTED_COUNTS.items():
        actual = counts[day]
        if actual != expected:
            issues.append({"type": "day_count_mismatch", "day": day, "expected": expected, "actual": actual})

    if len(rows) != sum(EXPECTED_COUNTS.values()):
        issues.append({"type": "total_count_mismatch", "expected": sum(EXPECTED_COUNTS.values()), "actual": len(rows)})

    for duplicated_id, count in id_counts.items():
        if count > 1:
            issues.append({"type": "duplicate_id", "id": duplicated_id, "count": count})

    for country, expected_id in DAY1_IDS.items():
        match = next((row for row in rows if row["day"] == 1 and row["country"] == country), None)
        if not match:
            issues.append({"type": "missing_day1_country", "country": country})
        elif match["id"] != expected_id:
            issues.append({"type": "day1_id_changed", "country": country, "expected": expected_id, "actual": match["id"]})

    by_day = defaultdict(list)
    for row in rows:
        by_day[row["day"]].append(row)
        if row["order"] < 1:
            issues.append({"type": "invalid_order", "id": row["id"], "order": row["order"]})
        if not row["country"] or not row["embassyName"] or not row["address"]:
            issues.append({"type": "missing_required_text", "id": row["id"]})
        if row["googleMapsQuery"]:
            expected_url = f"https://www.google.com/maps/search/?api=1&query={urllib.parse.quote(row['googleMapsQuery'])}"
            if not expected_url.startswith("https://www.google.com/maps/search/?api=1&query="):
                issues.append({"type": "invalid_generated_map_url", "id": row["id"]})
        else:
            issues.append({"type": "missing_google_maps_query", "id": row["id"], "country": row["country"], "sourceMap": row["sourceMap"]})

    for day, day_rows in by_day.items():
        orders = [row["order"] for row in sorted(day_rows, key=lambda item: item["order"])]
        expected_orders = list(range(1, len(day_rows) + 1))
        if orders != expected_orders:
            issues.append({"type": "order_sequence_mismatch", "day": day, "orders": orders})

    focus_terms = ["コートジボワール", "ボツワナ", "モーリタニア", "ハイチ", "ジブチ"]
    focus_rows = [
        {
            "id": row["id"],
            "day": row["day"],
            "order": row["order"],
            "country": row["country"],
            "embassyName": row["embassyName"],
            "address": row["address"],
            "hasGoogleMapsQuery": bool(row["googleMapsQuery"]),
        }
        for row in rows
        if row["country"] in focus_terms or any(term in row["embassyName"] for term in focus_terms)
    ]

    report = {
        "source": str(SOURCE_PATH),
        "total": len(rows),
        "expectedTotal": sum(EXPECTED_COUNTS.values()),
        "countsByDay": {f"Day{day}": counts[day] for day in range(1, 11)},
        "expectedCountsByDay": {f"Day{day}": EXPECTED_COUNTS[day] for day in range(1, 11)},
        "uniqueIdCount": len(set(ids)),
        "day1IdsPreserved": all(
            next((row for row in rows if row["day"] == 1 and row["country"] == country), {}).get("id") == expected_id
            for country, expected_id in DAY1_IDS.items()
        ),
        "validGoogleMapsCount": sum(1 for row in rows if row["googleMapsQuery"]),
        "missingGoogleMapsCount": sum(1 for row in rows if not row["googleMapsQuery"]),
        "focusRows": focus_rows,
        "issues": issues,
    }
    return report


def main():
    if not SOURCE_PATH.exists():
        print(f"Primary Source not found: {SOURCE_PATH}", file=sys.stderr)
        return 1

    rows, extraction_issues = extract_rows()
    assign_ids(rows)
    rows.sort(key=lambda item: (item["day"], item["order"]))
    write_data_js(rows)

    report = validate(rows, extraction_issues)
    report_path = ROOT / "docs" / "data-validation-v0.5.json"
    report_path.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")

    print(json.dumps(report, ensure_ascii=False, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
